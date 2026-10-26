from flask import Blueprint, jsonify, request
from datetime import datetime, date
from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel

from extensions import db
from models.asset import Asset
from models.vendor import Vendor
from models.asset_history import AssetHistory
from models.assignment import AssetAssignment
from routes.auth import roles_required


assets_bp = Blueprint(
    "assets",
    __name__,
    url_prefix="/api/assets"
)


@assets_bp.route("", methods=["GET"])
@roles_required(
    "SUPERUSER",
    "ADMIN",
    "IT"
)
def get_assets():

    search = request.args.get("search", "").strip()
    status = request.args.get("status", "").strip()
    category = request.args.get("category", "").strip()

    query = Asset.query

    if search:
        search_term = f"%{search}%"

        query = query.filter(
            db.or_(
                Asset.asset_tag.ilike(search_term),
                Asset.asset_type.ilike(search_term),
                Asset.category.ilike(search_term),
                Asset.brand.ilike(search_term),
                Asset.model.ilike(search_term),
                Asset.serial_number.ilike(search_term)
            )
        )

    if status and status != "ALL":
        query = query.filter(
            Asset.status == status
        )

    if category and category != "ALL":
        query = query.filter(
            Asset.category == category
        )

    assets = query.order_by(
        Asset.created_at.desc()
    ).all()

    return jsonify([
        asset_to_dict(asset)
        for asset in assets
    ])
    
    
# ============================================================
# GET DELETED ASSET HISTORY
# ============================================================

@assets_bp.route("/history", methods=["GET"])
@roles_required(
    "SUPERUSER",
    "ADMIN",
    "IT"
)
def get_asset_history():

    search = request.args.get(
        "search",
        ""
    ).strip()

    query = AssetHistory.query

    if search:

        search_term = f"%{search}%"

        query = query.filter(
            db.or_(
                AssetHistory.asset_tag.ilike(
                    search_term
                ),
                AssetHistory.asset_type.ilike(
                    search_term
                ),
                AssetHistory.category.ilike(
                    search_term
                ),
                AssetHistory.brand.ilike(
                    search_term
                ),
                AssetHistory.model.ilike(
                    search_term
                ),
                AssetHistory.serial_number.ilike(
                    search_term
                ),
                AssetHistory.vendor_name.ilike(
                    search_term
                ),
                AssetHistory.location.ilike(
                    search_term
                )
            )
        )

    history = query.order_by(
        AssetHistory.id.desc()
    ).all()

    return jsonify([
        asset_history_to_dict(item)
        for item in history
    ])


@assets_bp.route("/<int:asset_id>", methods=["GET"])
@roles_required(
    "SUPERUSER",
    "ADMIN",
    "IT"
)
def get_asset(asset_id):

    asset = db.session.get(
        Asset,
        asset_id
    )

    if not asset:
        return jsonify({
            "message": "Asset not found"
        }), 404

    return jsonify(
        asset_to_dict(asset)
    )


@assets_bp.route("", methods=["POST"])
@roles_required(
    "SUPERUSER",
    "ADMIN"
)
def create_asset():

    data = request.get_json() or {}

    required_fields = [
        "vendor_id",
        "asset_tag",
        "serial_number",
        "asset_type",
        "brand",
        "model",
        "received_date",
        "description",
        "status",
        "mode_of_transport",
        "warranty_status",
        "warranty_end",
        "department",
        "rental_cost",
        "location",
        "remarks"
    ]

    for field in required_fields:

        value = data.get(field)

        if value is None or str(value).strip() == "":
            return jsonify({
                "message":
                    f"{field.replace('_', ' ').title()} is required"
            }), 400

    existing_tag = Asset.query.filter_by(
        asset_tag=data["asset_tag"]
    ).first()

    if existing_tag:
        return jsonify({
            "message":
                "Asset tag already exists"
        }), 409

    serial_number = data.get("serial_number")

    if serial_number:

        existing_serial = Asset.query.filter_by(
            serial_number=serial_number
        ).first()

        if existing_serial:
            return jsonify({
                "message":
                    "Serial number already exists"
            }), 409

    vendor_id = data.get("vendor_id")

    if vendor_id:

        vendor = db.session.get(
            Vendor,
            vendor_id
        )

        if not vendor:
            return jsonify({
                "message":
                    "Selected vendor does not exist"
            }), 400

    asset = Asset(

        asset_tag=data["asset_tag"],

        asset_type=data["asset_type"],

        brand=data["brand"],

        model=data["model"],

        serial_number=data["serial_number"],

        vendor_id=vendor_id,

        received_date=parse_date(
            data["received_date"]
        ),

        item_specification=data.get(
            "item_specification"
        ),

        processor=data.get("processor"),

        ram=data.get("ram"),

        storage=data.get("storage"),

        display=data.get("display"),

        operating_system=data.get(
            "operating_system"
        ),

        description=data["description"],

        status=data["status"],

        mode_of_transport=data[
            "mode_of_transport"
        ],

        warranty_status=data[
            "warranty_status"
        ],

        warranty_end=parse_date(
            data["warranty_end"]
        ),

        department=data["department"],

        rental_cost=data["rental_cost"],

        location=data["location"],

        remarks=data["remarks"]
    )

    db.session.add(asset)
    db.session.commit()

    return jsonify({
        "message": "Asset created successfully",
        "asset": asset_to_dict(asset)
    }), 201


# ============================================================
# UPLOAD ASSETS FROM EXCEL
# ============================================================

@assets_bp.route(
    "/upload",
    methods=["POST"]
)
@roles_required(
    "SUPERUSER",
    "ADMIN"
)
def upload_assets():
    """
    Import assets from an Excel file.

    Rules:
    - Asset tag must not already exist.
    - Serial number must not already exist.
    - Vendor must already exist.
    - Duplicate rows are skipped.
    - Valid new assets are added.
    - Invalid rows are reported individually.
    """

    # --------------------------------------------------------
    # CHECK FILE
    # --------------------------------------------------------

    if "file" not in request.files:
        return jsonify({
            "success": False,
            "message": "No Excel file was uploaded."
        }), 400

    file = request.files["file"]

    if not file or not file.filename:
        return jsonify({
            "success": False,
            "message": "Please select an Excel file."
        }), 400

    filename = file.filename.lower()

    if not filename.endswith(
        (".xlsx", ".xls")
    ):
        return jsonify({
            "success": False,
            "message": (
                "Invalid file type. "
                "Please upload an Excel file."
            )
        }), 400

    # --------------------------------------------------------
    # LOAD EXCEL
    # --------------------------------------------------------

    try:

        workbook = load_workbook(
            file,
            data_only=True
        )

        worksheet = workbook.active

    except Exception as error:

        return jsonify({
            "success": False,
            "message": (
                "Unable to read the Excel file."
            ),
            "error": str(error)
        }), 400

    # --------------------------------------------------------
    # READ HEADERS
    # --------------------------------------------------------

    header_row = next(
        worksheet.iter_rows(
            min_row=1,
            max_row=1,
            values_only=True
        ),
        None
    )

    if not header_row:

        return jsonify({
            "success": False,
            "message": "The Excel file is empty."
        }), 400

    def normalize_header(value):

        if value is None:
            return ""

        return (
            str(value)
            .strip()
            .lower()
            .replace("_", " ")
            .replace("-", " ")
        )

    headers = {}

    for index, value in enumerate(
        header_row,
        start=1
    ):

        normalized = normalize_header(
            value
        )

        if normalized:
            headers[normalized] = index

    # --------------------------------------------------------
    # REQUIRED EXCEL COLUMNS
    # --------------------------------------------------------

    required_columns = [
        "asset tag",
        "serial number",
        "item type",
        "brand",
        "model",
        "vendor name",
        "date of arrived",
        "description",
        "asset status",
        "mode of transport",
        "warranty status",
        "warranty expiry date",
        "department",
        "rental cost",
        "location",
        "remarks"
    ]

    missing_columns = [
        column
        for column in required_columns
        if column not in headers
    ]

    if missing_columns:

        return jsonify({
            "success": False,
            "message": (
                "The Excel template is missing "
                "required columns."
            ),
            "missing_columns": missing_columns
        }), 400

    # --------------------------------------------------------
    # HELPER FUNCTIONS
    # --------------------------------------------------------

    def get_cell(row, column_name):

        column_index = headers.get(
            normalize_header(column_name)
        )

        if not column_index:
            return None

        return row[column_index - 1]

    def clean_value(value):

        if value is None:
            return ""

        return str(value).strip()

    def parse_excel_date(value):

        if value is None or value == "":
            return None

        # Excel datetime
        if isinstance(value, datetime):
            return value.date()

        # Python date
        if isinstance(value, date):
            return value

        # Excel serial date
        if isinstance(value, (int, float)):

            try:
                return from_excel(value).date()

            except Exception:
                pass

        value = str(value).strip()

        date_formats = [
            "%Y-%m-%d",
            "%d/%m/%Y",
            "%d-%m-%Y",
            "%m/%d/%Y"
        ]

        for date_format in date_formats:

            try:

                return datetime.strptime(
                    value,
                    date_format
                ).date()

            except ValueError:
                continue

        raise ValueError(
            "Invalid date. Use YYYY-MM-DD."
        )

    def parse_rental_cost(value):

        if value is None:
            return None

        if isinstance(
            value,
            (int, float)
        ):
            return float(value)

        value = str(value).strip()

        if not value:
            return None

        value = value.replace(
            ",",
            ""
        )

        try:

            return float(value)

        except ValueError:

            raise ValueError(
                "Rental Cost must be a valid number."
            )

    # --------------------------------------------------------
    # PROCESS ROWS
    # --------------------------------------------------------

    added = 0
    skipped = 0
    failed = 0

    errors = []
    added_assets = []

    rows = worksheet.iter_rows(
        min_row=2,
        values_only=True
    )

    for excel_row_number, row in enumerate(
        rows,
        start=2
    ):

        # ----------------------------------------------------
        # IGNORE COMPLETELY EMPTY ROWS
        # ----------------------------------------------------

        if not any(
            value is not None
            and str(value).strip() != ""
            for value in row
        ):
            continue

        try:

            # ------------------------------------------------
            # READ VALUES
            # ------------------------------------------------

            asset_tag = clean_value(
                get_cell(
                    row,
                    "Asset Tag"
                )
            )

            serial_number = clean_value(
                get_cell(
                    row,
                    "Serial Number"
                )
            )

            asset_type = clean_value(
                get_cell(
                    row,
                    "Item Type"
                )
            )

            brand = clean_value(
                get_cell(
                    row,
                    "Brand"
                )
            )

            model = clean_value(
                get_cell(
                    row,
                    "Model"
                )
            )

            vendor_name = clean_value(
                get_cell(
                    row,
                    "Vendor Name"
                )
            )

            received_date = parse_excel_date(
                get_cell(
                    row,
                    "Date of Arrived"
                )
            )

            processor = clean_value(
                get_cell(
                    row,
                    "Processor"
                )
            )

            ram = clean_value(
                get_cell(
                    row,
                    "RAM"
                )
            )

            storage = clean_value(
                get_cell(
                    row,
                    "Storage"
                )
            )

            display = clean_value(
                get_cell(
                    row,
                    "Display"
                )
            )

            operating_system = clean_value(
                get_cell(
                    row,
                    "OS"
                )
            )

            description = clean_value(
                get_cell(
                    row,
                    "Description"
                )
            )

            status = clean_value(
                get_cell(
                    row,
                    "Asset Status"
                )
            ).upper()

            mode_of_transport = clean_value(
                get_cell(
                    row,
                    "Mode of Transport"
                )
            )

            warranty_status = clean_value(
                get_cell(
                    row,
                    "Warranty Status"
                )
            )

            warranty_end = parse_excel_date(
                get_cell(
                    row,
                    "Warranty Expiry Date"
                )
            )

            department = clean_value(
                get_cell(
                    row,
                    "Department"
                )
            )

            rental_cost = parse_rental_cost(
                get_cell(
                    row,
                    "Rental Cost"
                )
            )

            location = clean_value(
                get_cell(
                    row,
                    "Location"
                )
            )

            remarks = clean_value(
                get_cell(
                    row,
                    "Remarks"
                )
            )

            # ------------------------------------------------
            # REQUIRED VALUE CHECK
            # ------------------------------------------------

            required_values = {
                "Asset Tag": asset_tag,
                "Serial Number": serial_number,
                "Item Type": asset_type,
                "Brand": brand,
                "Model": model,
                "Vendor Name": vendor_name,
                "Date of Arrived": received_date,
                "Description": description,
                "Asset Status": status,
                "Mode of Transport": mode_of_transport,
                "Warranty Status": warranty_status,
                "Warranty Expiry Date": warranty_end,
                "Department": department,
                "Location": location,
                "Remarks": remarks
            }

            missing_values = [
                field
                for field, value in required_values.items()
                if value is None
                or (
                    isinstance(value, str)
                    and not value.strip()
                )
            ]

            if missing_values:

                raise ValueError(
                    "Missing required values: "
                    + ", ".join(missing_values)
                )

            # ------------------------------------------------
            # VALIDATE STATUS
            # ------------------------------------------------

            allowed_statuses = {
                "AVAILABLE",
                "ASSIGNED",
                "REPAIR",
                "LOST",
                "RETIRED"
            }

            if status not in allowed_statuses:

                raise ValueError(
                    f"Invalid Asset Status '{status}'. "
                    "Allowed values: AVAILABLE, "
                    "ASSIGNED, REPAIR, LOST, RETIRED."
                )

            # ------------------------------------------------
            # CHECK DUPLICATE ASSET TAG
            # ------------------------------------------------

            existing_tag = Asset.query.filter_by(
                asset_tag=asset_tag
            ).first()

            if existing_tag:

                skipped += 1

                errors.append({
                    "row": excel_row_number,
                    "asset_tag": asset_tag,
                    "status": "SKIPPED",
                    "message": (
                        "Asset tag already exists."
                    )
                })

                continue

            # ------------------------------------------------
            # CHECK DUPLICATE SERIAL NUMBER
            # ------------------------------------------------

            existing_serial = Asset.query.filter_by(
                serial_number=serial_number
            ).first()

            if existing_serial:

                skipped += 1

                errors.append({
                    "row": excel_row_number,
                    "asset_tag": asset_tag,
                    "status": "SKIPPED",
                    "message": (
                        "Serial number already exists."
                    )
                })

                continue

            # ------------------------------------------------
            # FIND VENDOR
            # ------------------------------------------------

            vendor = Vendor.query.filter(
                db.func.lower(
                    Vendor.vendor_name
                ) == vendor_name.lower()
            ).first()

            if not vendor:

                raise ValueError(
                    f"Vendor '{vendor_name}' "
                    "does not exist."
                )

            # ------------------------------------------------
            # CREATE ITEM SPECIFICATION
            # ------------------------------------------------

            specification_parts = [
                processor,
                ram,
                storage,
                display,
                operating_system
            ]

            item_specification = " / ".join(
                value
                for value in specification_parts
                if value
            )

            # ------------------------------------------------
            # CREATE ASSET
            # ------------------------------------------------

            asset = Asset(

                asset_tag=asset_tag,

                asset_type=asset_type,

                brand=brand,

                model=model,

                serial_number=serial_number,

                vendor_id=vendor.id,

                received_date=received_date,

                item_specification=(
                    item_specification
                    or None
                ),

                processor=(
                    processor
                    or None
                ),

                ram=(
                    ram
                    or None
                ),

                storage=(
                    storage
                    or None
                ),

                display=(
                    display
                    or None
                ),

                operating_system=(
                    operating_system
                    or None
                ),

                description=description,

                status=status,

                mode_of_transport=(
                    mode_of_transport
                ),

                warranty_status=(
                    warranty_status
                ),

                warranty_end=warranty_end,

                department=department,

                rental_cost=rental_cost,

                location=location,

                remarks=remarks
            )

            savepoint = db.session.begin_nested()

            try:

                db.session.add(asset)

                db.session.flush()

                savepoint.commit()

                added += 1

                added_assets.append({
                    "row": excel_row_number,
                    "asset_tag": asset_tag,
                    "serial_number": serial_number
                })

            except Exception:
                savepoint.rollback()
                raise

        except Exception as error:

            failed += 1

            errors.append({
                "row": excel_row_number,
                "asset_tag": (
                    asset_tag
                    if "asset_tag" in locals()
                    else ""
                ),
                "status": "FAILED",
                "message": str(error)
            })

    # IMPORTANT:
    # Do NOT call db.session.rollback() here.
    #
    # A full rollback would remove assets that were
    # successfully processed in previous Excel rows.
    #
    # Each asset is already isolated using a nested
    # transaction/savepoint.

    # --------------------------------------------------------
    # COMMIT SUCCESSFULLY PROCESSED ASSETS
    # --------------------------------------------------------

    try:

        db.session.commit()

    except Exception as error:

        db.session.rollback()

        return jsonify({
            "success": False,
            "message": (
                "Failed to save imported assets."
            ),
            "error": str(error)
        }), 500

    # --------------------------------------------------------
    # FINAL RESPONSE
    # --------------------------------------------------------

    return jsonify({

        "success": True,

        "message": (
            "Excel asset import completed."
        ),

        "summary": {

            "added": added,

            "skipped": skipped,

            "failed": failed,

            "total_processed":
                added
                + skipped
                + failed
        },

        "added_assets": added_assets,

        "errors": errors

    }), 200
    

@assets_bp.route(
    "/<int:asset_id>",
    methods=["PUT"]
)
@roles_required(
    "SUPERUSER",
    "ADMIN"
)
def update_asset(asset_id):

    asset = db.session.get(
        Asset,
        asset_id
    )

    if not asset:
        return jsonify({
            "message": "Asset not found"
        }), 404

    data = request.get_json() or {}

    # -------------------------------------------------
    # ASSET TAG
    # -------------------------------------------------
    if "asset_tag" in data:

        asset_tag = data["asset_tag"].strip()

        existing = Asset.query.filter(
            Asset.asset_tag == asset_tag,
            Asset.id != asset.id
        ).first()

        if existing:
            return jsonify({
                "message": "Asset tag already exists"
            }), 409

        asset.asset_tag = asset_tag

    # -------------------------------------------------
    # SERIAL NUMBER
    # -------------------------------------------------
    if "serial_number" in data:

        serial_number = (
            data["serial_number"].strip()
            if data["serial_number"]
            else None
        )

        if serial_number:

            existing = Asset.query.filter(
                Asset.serial_number == serial_number,
                Asset.id != asset.id
            ).first()

            if existing:
                return jsonify({
                    "message": "Serial number already exists"
                }), 409

        asset.serial_number = serial_number

    # -------------------------------------------------
    # NORMAL TEXT FIELDS
    # -------------------------------------------------

    text_fields = [
        "asset_type",
        "category",
        "brand",
        "model",
        "item_specification",
        "processor",
        "ram",
        "storage",
        "display",
        "operating_system",
        "mode_of_transport",
        "warranty_status",
        "department",
        "location",
        "description",
        "remarks"
    ]

    for field in text_fields:

        if field in data:
            setattr(
                asset,
                field,
                data[field]
            )

    # -------------------------------------------------
    # STATUS
    # -------------------------------------------------

    if "status" in data:

        new_status = data["status"]

        allowed_statuses = {
            "AVAILABLE",
            "ASSIGNED",
            "REPAIR",
            "LOST",
            "RETIRED"
        }

        if new_status not in allowed_statuses:
            return jsonify({
                "message": "Invalid asset status"
            }), 400

        # Check whether asset currently has an
        # active assignment.
        active_assignment = None

        if hasattr(asset, "assignments"):

            active_assignment = next(
                (
                    assignment
                    for assignment in asset.assignments
                    if not assignment.actual_return_date
                ),
                None
            )

        # Assigned asset cannot be changed to AVAILABLE
        if active_assignment and new_status == "AVAILABLE":

            return jsonify({
                "message":
                    "This asset is currently assigned and cannot be marked as Available until it is returned."
            }), 400

        asset.status = new_status

    # -------------------------------------------------
    # VENDOR
    # -------------------------------------------------

    if "vendor_id" in data:

        vendor_id = data["vendor_id"]

        if vendor_id:

            vendor = db.session.get(
                Vendor,
                vendor_id
            )

            if not vendor:

                return jsonify({
                    "message":
                        "Selected vendor does not exist"
                }), 400

            asset.vendor_id = vendor_id

        else:

            asset.vendor_id = None

    # -------------------------------------------------
    # DATE FIELDS
    # -------------------------------------------------

    date_fields = [
        "purchase_date",
        "received_date",
        "warranty_start",
        "warranty_end"
    ]

    for field in date_fields:

        if field in data:

            setattr(
                asset,
                field,
                parse_date(data[field])
            )

    # -------------------------------------------------
    # RENTAL COST
    # -------------------------------------------------

    if "rental_cost" in data:

        value = data["rental_cost"]

        if value in ("", None):
            asset.rental_cost = None
        else:
            try:
                asset.rental_cost = float(value)
            except (ValueError, TypeError):

                return jsonify({
                    "message":
                        "Rental cost must be a valid number"
                }), 400

    # -------------------------------------------------
    # COMMIT
    # -------------------------------------------------

    try:

        db.session.commit()

    except Exception as error:

        db.session.rollback()

        print(
            "Asset update error:",
            error
        )

        return jsonify({
            "message":
                "Failed to update asset"
        }), 500

    return jsonify({
        "message":
            "Asset updated successfully",
        "asset":
            asset_to_dict(asset)
    }), 200


# ============================================================
# DELETE ASSET
# ============================================================

@assets_bp.route(
    "/<int:asset_id>",
    methods=["DELETE"]
)
@roles_required(
    "SUPERUSER",
    "ADMIN"
)
def delete_asset(asset_id):

    asset = db.session.get(
        Asset,
        asset_id
    )

    if not asset:
        return jsonify({
            "message": "Asset not found"
        }), 404

    # --------------------------------------------------------
    # PROTECT ASSIGNED ASSETS
    # --------------------------------------------------------

    active_assignment = AssetAssignment.query.filter(
        AssetAssignment.asset_id == asset.id,
        AssetAssignment.actual_return_date.is_(None)
    ).first()

    # Block deletion if:
    # 1. There is an active assignment
    # OR
    # 2. Asset status is ASSIGNED

    if active_assignment or asset.status == "ASSIGNED":

        return jsonify({
            "message": (
                "This asset cannot be deleted because it is "
                "currently assigned. Return the asset first."
            )
        }), 409

    # --------------------------------------------------------
    # SAVE ASSET SNAPSHOT INTO HISTORY
    # --------------------------------------------------------

    history = AssetHistory(
        original_asset_id=asset.id,
        asset_tag=asset.asset_tag,
        asset_type=asset.asset_type,
        category=asset.category,
        brand=asset.brand,
        model=asset.model,
        serial_number=asset.serial_number,
        vendor_id=asset.vendor_id,

        vendor_name=(
            asset.vendor.vendor_name
            if asset.vendor
            else None
        ),

        purchase_date=asset.purchase_date,
        received_date=asset.received_date,
        purchase_cost=asset.purchase_cost,
        warranty_start=asset.warranty_start,
        warranty_end=asset.warranty_end,
        status=asset.status,
        location=asset.location,
        description=asset.description
    )

    db.session.add(history)

    # --------------------------------------------------------
    # DELETE ASSET
    # --------------------------------------------------------

    try:

        db.session.delete(asset)

        db.session.commit()

        return jsonify({
            "message": (
                "Asset deleted successfully "
                "and moved to history"
            )
        }), 200

    except Exception as error:

        db.session.rollback()

        return jsonify({
            "message": "Failed to delete asset",
            "error": str(error)
        }), 500


@assets_bp.route(
    "/<int:asset_id>/status",
    methods=["PUT"]
)
@roles_required(
    "SUPERUSER",
    "ADMIN"
)
def update_asset_status(asset_id):

    asset = db.session.get(
        Asset,
        asset_id
    )

    if not asset:
        return jsonify({
            "message": "Asset not found"
        }), 404

    data = request.get_json() or {}

    new_status = data.get("status")

    allowed_statuses = {
        "AVAILABLE",
        "ASSIGNED",
        "REPAIR",
        "LOST",
        "RETIRED"
    }

    if new_status not in allowed_statuses:
        return jsonify({
            "message":
                "Invalid asset status"
        }), 400

    asset.status = new_status

    db.session.commit()

    return jsonify({
        "message":
            "Asset status updated successfully",
        "asset":
            asset_to_dict(asset)
    })


def parse_date(value):

    if not value:
        return None

    if hasattr(value, "year"):
        return value

    try:
        return datetime.strptime(
            value,
            "%Y-%m-%d"
        ).date()

    except ValueError:
        raise ValueError(
            "Date must be in YYYY-MM-DD format"
        )


def asset_history_to_dict(history):

    return {
        "id": history.id,

        "original_asset_id":
            history.original_asset_id,

        "asset_tag":
            history.asset_tag,

        "asset_type":
            history.asset_type,

        "category":
            history.category,

        "brand":
            history.brand,

        "model":
            history.model,

        "serial_number":
            history.serial_number,

        "vendor_id":
            history.vendor_id,

        "vendor_name":
            history.vendor_name,

        "purchase_date":
            history.purchase_date.isoformat()
            if history.purchase_date
            else None,

        "received_date":
            history.received_date.isoformat()
            if history.received_date
            else None,

        "purchase_cost":
            float(history.purchase_cost)
            if history.purchase_cost is not None
            else None,

        "warranty_start":
            history.warranty_start.isoformat()
            if history.warranty_start
            else None,

        "warranty_end":
            history.warranty_end.isoformat()
            if history.warranty_end
            else None,

        "status":
            history.status,

        "location":
            history.location,

        "description":
            history.description,

        "created_at":
            history.created_at.isoformat()
            if getattr(history, "created_at", None)
            else None,
            
        "deleted_at": (
            history.deleted_at.strftime("%d/%m/%Y")
            if history.deleted_at
            else None
        ),
    }
    
    
# ============================================================
# DELETE ASSET HISTORY RECORD
# ============================================================

@assets_bp.route(
    "/history/<int:history_id>",
    methods=["DELETE"]
)
@roles_required(
    "SUPERUSER",
    "ADMIN"
)
def delete_asset_history(history_id):

    history = db.session.get(
        AssetHistory,
        history_id
    )

    if not history:

        return jsonify({
            "message":
                "Asset history record not found"
        }), 404

    try:

        db.session.delete(history)

        db.session.commit()

        return jsonify({
            "message":
                "Asset history record deleted successfully"
        }), 200

    except Exception as error:

        db.session.rollback()

        return jsonify({
            "message":
                "Failed to delete asset history record",
            "error":
                str(error)
        }), 500


def asset_to_dict(asset):

    return {
        "id": asset.id,

        "asset_tag": asset.asset_tag,

        "asset_type": asset.asset_type,

        "category": asset.category,

        "brand": asset.brand,

        "model": asset.model,

        "serial_number": asset.serial_number,

        "vendor_id": asset.vendor_id,

        "vendor_name":
            asset.vendor.vendor_name
            if asset.vendor
            else None,

        "received_date":
            asset.received_date.isoformat()
            if asset.received_date
            else None,

        "item_specification":
            asset.item_specification,

        "processor":
            asset.processor,

        "ram":
            asset.ram,

        "storage":
            asset.storage,

        "display":
            asset.display,

        "operating_system":
            asset.operating_system,

        "description":
            asset.description,

        "status":
            asset.status,

        "mode_of_transport":
            asset.mode_of_transport,

        "warranty_status":
            asset.warranty_status,

        "warranty_end":
            asset.warranty_end.isoformat()
            if asset.warranty_end
            else None,

        "department":
            asset.department,

        "rental_cost":
            float(asset.rental_cost)
            if asset.rental_cost is not None
            else None,

        "location":
            asset.location,

        "remarks":
            asset.remarks,

        "created_at":
            asset.created_at.isoformat()
            if asset.created_at
            else None,

        "updated_at":
            asset.updated_at.isoformat()
            if asset.updated_at
            else None
    }